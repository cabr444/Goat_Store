"""Build the editable English and Spanish Goats Store manuals.

Requires python-docx. Run from the repository root:
  python scripts/build-documentation.py --output-dir work/documentation
"""
from __future__ import annotations
import argparse
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

PURPLE = '7041D9'
INK = '211A35'
MUTED = '615C70'
LIGHT = 'F4F1FA'
CONFIG = {
    'en': {
        'source': 'goat-store-en.md', 'file': 'Goat_Store_Documentation_EN.docx',
        'subtitle': 'Functional, Technical & Operational Documentation',
        'tagline': 'Streaming commerce • Administration • Deployment • Maintenance',
        'edition': 'ENGLISH EDITION', 'contents': 'Contents',
        'guide': 'Select a section to navigate. All content, tables and procedures are editable.',
        'summary': 'A complete reference for operating, administering and maintaining Goats Store. Based on the implemented application, with clear separation between current capabilities, operating procedures and future improvements.',
        'version': 'Document version 2.0', 'date': '26 October 2026',
        'baseline': 'Application 2.0.0  |  Source 22cd272  |  password-reset',
        'footer': 'GOATS STORE  •  Documentation 2.0', 'page': 'Page', 'of': 'of',
        'lang': 'en-US',
    },
    'es': {
        'source': 'goat-store-es.md', 'file': 'Goat_Store_Documentacion_ES.docx',
        'subtitle': 'Documentación Funcional, Técnica y Operativa',
        'tagline': 'Comercio de streaming • Administración • Despliegue • Mantenimiento',
        'edition': 'EDICIÓN EN ESPAÑOL', 'contents': 'Contenido',
        'guide': 'Selecciona una sección para navegar. Todo el contenido, tablas y procedimientos son editables.',
        'summary': 'Referencia completa para operar, administrar y mantener Goats Store. Basada en la aplicación implementada, con una distinción clara entre capacidades actuales, procedimientos operativos y mejoras futuras.',
        'version': 'Versión documental 2.0', 'date': '26 de octubre de 2026',
        'baseline': 'Aplicación 2.0.0  |  Código 22cd272  |  password-reset',
        'footer': 'GOATS STORE  •  Documentación 2.0', 'page': 'Página', 'of': 'de',
        'lang': 'es-CO',
    },
}


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:shd'); node.set(qn('w:fill'), color); props.append(node)


def set_cell_margins(cell):
    props = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
    for edge, value in [('top', 85), ('bottom', 85), ('left', 100), ('right', 100)]:
        node = OxmlElement('w:' + edge); node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa'); margins.append(node)
    props.append(margins)


def run_style(run, size=10.5, color=INK, bold=False, font='Calibri'):
    run.font.name = font; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold


def link(paragraph, text, anchor):
    node = OxmlElement('w:hyperlink'); node.set(qn('w:anchor'), anchor); node.set(qn('w:history'), '1')
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), INK); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '21'); props.append(size)
    run.append(props); value = OxmlElement('w:t'); value.text = text; run.append(value); node.append(run); paragraph._p.append(node)


def bookmark(paragraph, identifier):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(identifier)); start.set(qn('w:name'), f'section_{identifier}')
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(identifier))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def field(paragraph, instruction):
    node = OxmlElement('w:fldSimple'); node.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r'); value = OxmlElement('w:t'); value.text = '1'; run.append(value); node.append(run); paragraph._p.append(node)


def styles(document, language):
    normal = document.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), language); normal.element.get_or_add_rPr().append(lang)
    for name, size, color in [('Title', 39, INK), ('Subtitle', 16, MUTED), ('Heading 1', 22, PURPLE), ('Heading 2', 14, INK)]:
        style = document.styles[name]; style.font.name = 'Calibri'; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith('Heading')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(12)
    document.styles['Heading 1'].paragraph_format.page_break_before = True
    for name in ['List Bullet', 'List Number']:
        style = document.styles[name]; style.font.name = 'Calibri'; style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(5); style.paragraph_format.line_spacing = 1.1
    code = document.styles.add_style('Goat Code', 1)
    code.font.name = 'Consolas'; code.font.size = Pt(9); code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.line_spacing = 1.1; code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.left_indent = Cm(.25); code.paragraph_format.right_indent = Cm(.15)
    code.paragraph_format.keep_together = True


def add_cover(document, config):
    p = document.add_paragraph(); p.paragraph_format.space_before = Pt(48)
    run_style(p.add_run('Goats'), 45, INK, True); run_style(p.add_run('.'), 45, PURPLE, True)
    p = document.add_paragraph(config['edition']); p.paragraph_format.space_before = Pt(38)
    for run in p.runs: run_style(run, 10, PURPLE, True)
    p = document.add_paragraph('GOATS STORE', 'Title'); p.paragraph_format.space_after = Pt(14)
    p = document.add_paragraph(config['subtitle'], 'Subtitle'); p.paragraph_format.space_after = Pt(16)
    p = document.add_paragraph(config['tagline']); p.paragraph_format.space_after = Pt(25)
    for run in p.runs: run_style(run, 10.5, MUTED)
    p = document.add_paragraph(config['summary']); p.paragraph_format.space_after = Pt(30)
    p = document.add_paragraph(config['version'] + '\n' + config['date'])
    for run in p.runs: run_style(run, 11, INK, True)
    p = document.add_paragraph(config['baseline']); p.paragraph_format.space_before = Pt(12)
    for run in p.runs: run_style(run, 9, MUTED)
    p = document.add_paragraph('WORD / DOCX  •  EDITABLE'); p.paragraph_format.space_before = Pt(12)
    for run in p.runs: run_style(run, 9, PURPLE, True)
    document.add_page_break()


def add_contents(document, headings, config):
    p = document.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    run_style(p.add_run(config['contents']), 24, PURPLE, True)
    p = document.add_paragraph(config['guide']); p.paragraph_format.space_after = Pt(15)
    for run in p.runs: run_style(run, 9.5, MUTED)
    for i, title in enumerate(headings, 1):
        p = document.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.05
        link(p, title, f'section_{i}')


def table(document, rows):
    headers = rows[0]; columns = len(headers)
    result = document.add_table(rows=1, cols=columns)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER; result.autofit = False
    if columns == 5: fractions = [.44, .17, .18, .11, .10]
    elif columns == 4: fractions = [.46, .17, .21, .16]
    elif columns == 3:
        if headers[0] in ['Model', 'Modelo']: fractions = [.20, .55, .25]
        elif 'endpoint' in headers[0].lower(): fractions = [.43, .19, .38]
        else: fractions = [.32, .36, .32]
    elif columns == 2: fractions = [.35, .65]
    else: fractions = [1 / columns] * columns
    for col, fraction in zip(result.columns, fractions): col.width = Cm(17 * fraction)
    for i, values in enumerate(rows):
        cells = result.rows[0].cells if i == 0 else result.add_row().cells
        for cell, value, fraction in zip(cells, values, fractions):
            cell.width = Cm(17 * fraction); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = i == 0
            run_style(p.add_run(value), 8.5 if columns > 2 else 9.2, 'FFFFFF' if i == 0 else INK, i == 0)
            if i == 0: shade(cell, PURPLE)
            elif i % 2: shade(cell, LIGHT)
        props = result.rows[i]._tr.get_or_add_trPr()
        keep = OxmlElement('w:cantSplit'); props.append(keep)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); props.append(repeat)
    p = document.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(2); p.add_run().font.size = Pt(2)


def new_numbered_list(document):
    root = document.part.numbering_part.element
    base_id = document.styles['List Number'].element.pPr.numPr.numId.val
    base = next(node for node in root if node.tag == qn('w:num') and node.get(qn('w:numId')) == str(base_id))
    abstract_id = base.find(qn('w:abstractNumId')).get(qn('w:val'))
    next_id = max(int(node.get(qn('w:numId'))) for node in root if node.tag == qn('w:num')) + 1
    number = OxmlElement('w:num'); number.set(qn('w:numId'), str(next_id))
    abstract = OxmlElement('w:abstractNumId'); abstract.set(qn('w:val'), abstract_id); number.append(abstract)
    override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:startOverride'); start.set(qn('w:val'), '1'); override.append(start); number.append(override)
    root.append(number)
    return next_id


def build(source, destination, config):
    text = source.read_text(encoding='utf-8')
    headings = [line[3:].strip() for line in text.splitlines() if line.startswith('## ')]
    document = Document(); section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2); section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2); section.right_margin = Cm(2)
    section.header_distance = Cm(.8); section.footer_distance = Cm(.8)
    section.different_first_page_header_footer = True
    styles(document, config['lang'])
    document.core_properties.title = 'Goats Store — ' + config['subtitle']
    document.core_properties.subject = config['subtitle']
    document.core_properties.author = 'Goats Store'
    document.core_properties.keywords = 'Goats Store, streaming, administration, operations, documentation'
    document.core_properties.created = datetime.now(timezone.utc)
    document.core_properties.modified = datetime.now(timezone.utc)
    document.core_properties.revision = 2
    document.core_properties.comments = 'Editable documentation. Source baseline: 22cd272; documentation 2.0.'
    p = section.header.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    run_style(p.add_run('Goats.'), 11, PURPLE, True)
    run_style(p.add_run('   /   ' + config['edition']), 8, MUTED)
    p = section.footer.paragraphs[0]; p.paragraph_format.tab_stops.add_tab_stop(Cm(12.1))
    run_style(p.add_run(config['footer'] + '\t' + config['page'] + ' '), 8, MUTED)
    field(p, 'PAGE'); run_style(p.add_run(' ' + config['of'] + ' '), 8, MUTED); field(p, 'NUMPAGES')
    add_cover(document, config); add_contents(document, headings, config)
    lines = text.splitlines(); index = 0; heading_number = 0; current_list = None
    while index < len(lines):
        line = lines[index]
        if not line.strip() or line.startswith('# '): index += 1; continue
        if line.startswith('## '):
            heading_number += 1
            p = document.add_paragraph(line[3:].strip(), 'Heading 1'); bookmark(p, heading_number); index += 1; continue
        if line.startswith('### '): document.add_paragraph(line[4:].strip(), 'Heading 2'); index += 1; continue
        if line.startswith('```'):
            code = []; index += 1
            while index < len(lines) and not lines[index].startswith('```'): code.append(lines[index]); index += 1
            p = document.add_paragraph('\n'.join(code), 'Goat Code')
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), LIGHT); p._p.get_or_add_pPr().append(shd)
            index += 1; continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                values = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', cell) for cell in values): rows.append(values)
                index += 1
            table(document, rows); continue
        if re.match(r'^\d+\. ', line):
            if line.startswith('1. ') or current_list is None: current_list = new_numbered_list(document)
            p = document.add_paragraph(re.sub(r'^\d+\. ', '', line), 'List Number')
            numbering = p._p.get_or_add_pPr().get_or_add_numPr()
            numbering.get_or_add_ilvl().val = 0; numbering.get_or_add_numId().val = current_list
            index += 1; continue
        if line.startswith('- '): document.add_paragraph(line[2:], 'List Bullet'); index += 1; continue
        paragraph = [line.strip()]; index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r'^(#|\||```|\d+\. |\- )', lines[index]):
            paragraph.append(lines[index].strip()); index += 1
        p = document.add_paragraph(' '.join(paragraph))
        if p.text.startswith(('Glossary:', 'Glosario:')): p.paragraph_format.keep_together = True
    destination.parent.mkdir(parents=True, exist_ok=True); document.save(destination)
    print(f'{destination.name}: {heading_number} sections; editable DOCX created.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=Path(__file__).resolve().parent.parent / 'docs')
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    for config in CONFIG.values(): build(args.source_dir / config['source'], args.output_dir / config['file'], config)
    archive = args.output_dir / 'Goat_Store_Documentation_Sources.zip'
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as output:
        for config in CONFIG.values():
            output.write(args.source_dir / config['source'], 'docs/' + config['source'])
        readme = args.source_dir / 'README.md'
        if readme.exists(): output.write(readme, 'docs/README.md')
        output.write(Path(__file__), 'scripts/build-documentation.py')
    print(f'{archive.name}: bilingual sources and generator packaged.')


if __name__ == '__main__': main()
